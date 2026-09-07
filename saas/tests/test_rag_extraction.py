import io

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from src.rag.extraction import ExtractionError, extract_text, sha256_of


def test_sha256_of_is_deterministic_and_matches_hashlib():
    import hashlib

    data = "テストデータ".encode("utf-8")
    assert sha256_of(data) == hashlib.sha256(data).hexdigest()


def test_extract_text_plain_txt():
    data = "これはテキストファイルの中身です。".encode("utf-8")
    assert extract_text(data, "text/plain") == "これはテキストファイルの中身です。"


def test_extract_text_markdown():
    data = "# 見出し\n本文です。".encode("utf-8")
    assert extract_text(data, "text/markdown") == "# 見出し\n本文です。"


def test_extract_text_rejects_unsupported_mime_type():
    with pytest.raises(ExtractionError):
        extract_text(b"data", "image/png")


def test_extract_text_rejects_undecodable_bytes_for_plain_text():
    with pytest.raises(ExtractionError):
        extract_text(b"\xff\xfe\x00\x01", "text/plain")


def test_extract_text_rejects_empty_content_after_strip():
    with pytest.raises(ExtractionError):
        extract_text("   \n  ".encode("utf-8"), "text/plain")


def test_extract_text_docx_reads_paragraphs():
    buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_paragraph("一段落目です。")
    doc.add_paragraph("二段落目です。")
    doc.save(buf)

    text = extract_text(
        buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert "一段落目です。" in text
    assert "二段落目です。" in text


def test_extract_text_pdf_with_no_text_raises_extraction_error():
    # テキストを含まない(画像のみ相当の)空白PDFは、スキャンPDFと同様に扱われる
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)

    with pytest.raises(ExtractionError):
        extract_text(buf.getvalue(), "application/pdf")


def test_extract_text_pdf_invalid_bytes_raises_extraction_error():
    with pytest.raises(ExtractionError):
        extract_text(b"not a real pdf", "application/pdf")
